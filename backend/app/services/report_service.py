from datetime import date
from pathlib import Path

from sqlalchemy.orm import Session

from app.models.task import Task

from openpyxl import Workbook

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib import colors


# =========================================================
# REPORT DIRECTORY
# =========================================================

REPORT_DIR = Path("generated_reports")
REPORT_DIR.mkdir(parents=True, exist_ok=True)


# =========================================================
# STATUS DEFINITIONS
# =========================================================

# The project has used both status naming conventions:
#
#   Pending / In Progress / Completed
#
# and:
#
#   To Do / In Progress / Done
#
# Reports support both so existing data is not broken.

COMPLETED_STATUSES = ["Completed", "Done"]

PENDING_STATUSES = ["Pending", "To Do"]

IN_PROGRESS_STATUSES = ["In Progress"]


# =========================================================
# DAILY REPORT
# =========================================================

def generate_daily_report(db: Session):

    # -----------------------------------------------------
    # Total active tasks
    # -----------------------------------------------------

    total_tasks = (
        db.query(Task)
        .filter(
            Task.is_deleted == False
        )
        .count()
    )

    # -----------------------------------------------------
    # Completed tasks
    # -----------------------------------------------------

    completed = (
        db.query(Task)
        .filter(
            Task.status.in_(COMPLETED_STATUSES),
            Task.is_deleted == False
        )
        .count()
    )

    # -----------------------------------------------------
    # Pending / To Do tasks
    # -----------------------------------------------------

    pending = (
        db.query(Task)
        .filter(
            Task.status.in_(PENDING_STATUSES),
            Task.is_deleted == False
        )
        .count()
    )

    # -----------------------------------------------------
    # In-progress tasks
    # -----------------------------------------------------

    in_progress = (
        db.query(Task)
        .filter(
            Task.status.in_(IN_PROGRESS_STATUSES),
            Task.is_deleted == False
        )
        .count()
    )

    # -----------------------------------------------------
    # Overdue tasks
    #
    # A task is overdue when:
    # - it has a due date
    # - due date is before today
    # - it is not completed
    # - it is not deleted
    # -----------------------------------------------------

    overdue = (
        db.query(Task)
        .filter(
            Task.due_date.isnot(None),
            Task.due_date < date.today(),
            ~Task.status.in_(COMPLETED_STATUSES),
            Task.is_deleted == False
        )
        .count()
    )

    # -----------------------------------------------------
    # Response
    # -----------------------------------------------------

    return {
        "report_type": "Daily",
        "generated_on": str(date.today()),
        "summary": {
            "total_tasks": total_tasks,
            "completed": completed,
            "pending": pending,
            "in_progress": in_progress,
            "overdue": overdue,
        },
    }


# =========================================================
# WEEKLY REPORT
# =========================================================

def generate_weekly_report(db: Session):

    # Currently the project dashboard/reporting requirement
    # uses the same task summary structure for the weekly
    # report.

    daily = generate_daily_report(db)

    return {
        "report_type": "Weekly",
        "generated_on": str(date.today()),
        "summary": daily["summary"],
    }


# =========================================================
# REPORT HISTORY
# =========================================================

def report_history():

    return [
        {
            "report_id": 1,
            "type": "Daily",
            "generated_on": "2026-07-26",
        },
        {
            "report_id": 2,
            "type": "Weekly",
            "generated_on": "2026-07-21",
        },
    ]


# =========================================================
# CREATE EXCEL REPORT
# =========================================================

def create_excel_report(report_data, filename):

    file_path = REPORT_DIR / filename

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Task Report"

    # -----------------------------------------------------
    # Header
    # -----------------------------------------------------

    sheet["A1"] = "AI Task Management System"
    sheet["A2"] = report_data["report_type"] + " Report"
    sheet["A3"] = "Generated On"
    sheet["B3"] = report_data["generated_on"]

    # -----------------------------------------------------
    # Table header
    # -----------------------------------------------------

    sheet["A5"] = "Metric"
    sheet["B5"] = "Count"

    summary = report_data["summary"]

    rows = [
        ("Total Tasks", summary["total_tasks"]),
        ("Completed", summary["completed"]),
        ("Pending", summary["pending"]),
        ("In Progress", summary["in_progress"]),
        ("Overdue", summary["overdue"]),
    ]

    row_number = 6

    for metric, count in rows:

        sheet.cell(
            row=row_number,
            column=1,
            value=metric,
        )

        sheet.cell(
            row=row_number,
            column=2,
            value=count,
        )

        row_number += 1

    # -----------------------------------------------------
    # Column widths
    # -----------------------------------------------------

    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 15

    workbook.save(file_path)

    return file_path


# =========================================================
# CREATE PDF REPORT
# =========================================================

def create_pdf_report(report_data, filename):

    file_path = REPORT_DIR / filename

    document = SimpleDocTemplate(
        str(file_path),
        pagesize=A4,
    )

    summary = report_data["summary"]

    data = [
        ["Metric", "Count"],
        ["Total Tasks", summary["total_tasks"]],
        ["Completed", summary["completed"]],
        ["Pending", summary["pending"]],
        ["In Progress", summary["in_progress"]],
        ["Overdue", summary["overdue"]],
    ]

    table = Table(data)

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.grey,
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    1,
                    colors.black,
                ),
                (
                    "ALIGN",
                    (1, 1),
                    (-1, -1),
                    "CENTER",
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
            ]
        )
    )

    document.build([table])

    return file_path


# =========================================================
# DOWNLOAD REPORT
# =========================================================

def download_report(db: Session, report_type: str):

    report_type = report_type.strip().lower()

    # -----------------------------------------------------
    # Validate report type
    # -----------------------------------------------------

    if report_type not in ["excel", "pdf"]:

        return {
            "message": "Invalid report type",
            "supported_types": [
                "excel",
                "pdf",
            ],
        }

    # -----------------------------------------------------
    # Generate report data
    # -----------------------------------------------------

    report_data = generate_daily_report(db)

    # -----------------------------------------------------
    # Excel
    # -----------------------------------------------------

    if report_type == "excel":

        filename = "daily_report.xlsx"

        return create_excel_report(
            report_data,
            filename,
        )

    # -----------------------------------------------------
    # PDF
    # -----------------------------------------------------

    filename = "daily_report.pdf"

    return create_pdf_report(
        report_data,
        filename,
    )